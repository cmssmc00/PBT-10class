#!/usr/bin/env python3
"""Build the age/location sidecar used by the hierarchical prior experiment."""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pathology_credit.data import LOCATION_NAMES, LOCATION_TO_INDEX


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=PROJECT_ROOT / "Dell Children's Brain tumor Database.xlsx",
    )
    parser.add_argument(
        "--patients-json",
        type=Path,
        default=PROJECT_ROOT / "data" / "patients_206.json",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "data" / "clinical_206.json",
    )
    return parser.parse_args()


def normalise_location(value: Any) -> tuple[str, bool]:
    """Map free-text anatomy to a label-only-independent location vocabulary."""

    if value is None or not str(value).strip():
        return "Unknown", False
    text = re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()
    if not text or text in {"na", "n a", "nd", "n d", "unknown", "unspecified"}:
        return "Unknown", False

    # Specific structures precede broad compartment words when both occur.
    rules = (
        ("Pineal region", ("pineal",)),
        (
            "Sellar/suprasellar/optic",
            ("sellar", "suprasellar", "pituitary", "optic", "chiasm"),
        ),
        ("Spinal cord", ("spinal",)),
        (
            "Cerebellopontine angle",
            (
                "cerebellopontine",
                "cerebellar pontine",
                "cerebropontine",
                "cerebropontile",
                "cp angle",
                "c p angle",
                "cpa",
            ),
        ),
        (
            "Brainstem/tectum",
            ("brainstem", "brain stem", "pons", "pontine", "medulla", "tectum", "tectal"),
        ),
        (
            "Ventricular/periventricular",
            ("ventricle", "ventricular", "venticular", "periventricular"),
        ),
        ("Cerebellum", ("cerebell",)),
        ("Posterior fossa, NOS", ("posterior fossa",)),
        (
            "Thalamus/deep gray",
            ("thalam", "basal ganglia", "deep gray", "deep grey"),
        ),
        (
            "Cerebral hemispheric/lobar",
            (
                "frontal",
                "temporal",
                "parietal",
                "occipital",
                "insular",
                "hippocamp",
                "cingulate",
                "corpus callosum",
                "splenium",
                "cerebral",
                "supratentorial",
                "hemisphere",
            ),
        ),
    )
    for group, fragments in rules:
        if any(fragment in text for fragment in fragments):
            return group, True
    return "Unknown", False


def _finite_age(value: Any, case_id: str) -> tuple[float | None, bool]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, False
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise TypeError(f"{case_id}: age must be numeric, got {value!r}")
    age = float(value)
    if not math.isfinite(age) or not 0.0 <= age <= 120.0:
        raise ValueError(f"{case_id}: age must be finite and in [0, 120]")
    return age, True


def _resolve(path: Path) -> Path:
    return path if path.is_absolute() else PROJECT_ROOT / path


def main() -> None:
    args = parse_args()
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "openpyxl is needed only to rebuild the clinical JSON sidecar"
        ) from exc

    workbook_path = _resolve(args.workbook)
    patients_path = _resolve(args.patients_json)
    output_path = _resolve(args.output)
    patients = json.loads(patients_path.read_text(encoding="utf-8"))
    if not isinstance(patients, list):
        raise TypeError("patients JSON must be a list")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["coded sheet"]
        headers = {
            str(cell.value).strip(): column
            for column, cell in enumerate(worksheet[1], start=1)
            if cell.value is not None
        }
        required = {"Code", "Patient age", "Tumor location", "10 Classications"}
        if not required.issubset(headers):
            raise ValueError(f"Workbook is missing headers: {sorted(required - set(headers))}")

        clinical_rows: list[dict[str, Any]] = []
        seen: set[str] = set()
        for patient in patients:
            if not isinstance(patient, dict):
                raise TypeError("Each patient JSON entry must be an object")
            case_id = patient.get("case_id")
            source_row = patient.get("source_excel_row")
            if not isinstance(case_id, str) or not case_id:
                raise ValueError("Patient JSON contains an invalid case_id")
            if isinstance(source_row, bool) or not isinstance(source_row, int) or source_row < 2:
                raise ValueError(f"{case_id}: invalid source_excel_row")
            if case_id in seen:
                raise ValueError(f"Duplicate patient case_id: {case_id}")
            seen.add(case_id)

            workbook_case = worksheet.cell(source_row, headers["Code"]).value
            workbook_class = worksheet.cell(
                source_row, headers["10 Classications"]
            ).value
            if str(workbook_case).strip() != case_id:
                raise ValueError(
                    f"{case_id}: workbook row {source_row} has code {workbook_case!r}"
                )
            if workbook_class != patient.get("classification"):
                raise ValueError(f"{case_id}: workbook/patient class mismatch")

            raw_age = worksheet.cell(source_row, headers["Patient age"]).value
            age, age_mask = _finite_age(raw_age, case_id)
            raw_location = worksheet.cell(
                source_row, headers["Tumor location"]
            ).value
            location_group, location_mask = normalise_location(raw_location)
            clinical_rows.append(
                {
                    "case_id": case_id,
                    "source_excel_row": source_row,
                    "age": age,
                    "age_mask": age_mask,
                    "location_raw": raw_location,
                    "location_group": location_group,
                    "location_index": LOCATION_TO_INDEX[location_group],
                    "location_mask": location_mask,
                }
            )
    finally:
        workbook.close()

    payload = {
        "schema_version": 1,
        "age_unit": "years",
        "location_vocabulary": list(LOCATION_NAMES),
        "normalization_note": (
            "Location mapping uses anatomy terms only and never reads the tumor label. "
            "Unspecified, blank, and unmatched values map to Unknown."
        ),
        "patients": clinical_rows,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    os.replace(temporary, output_path)

    location_counts = Counter(row["location_group"] for row in clinical_rows)
    print(f"Wrote {len(clinical_rows)} patients to {output_path}")
    print(f"Age available: {sum(row['age_mask'] for row in clinical_rows)}")
    for location in LOCATION_NAMES:
        print(f"  {location}: {location_counts[location]}")


if __name__ == "__main__":
    main()
